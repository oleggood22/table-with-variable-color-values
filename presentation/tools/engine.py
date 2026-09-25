"""Tiny layout engine: one slide description -> PPTX (python-pptx) and HTML.

Every element is placed on a 1920x1080 px canvas. The same numbers drive both
outputs, so the PowerPoint file and the HTML copy stay in sync.
"""
import copy
import html
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

W, H = 1920, 1080
EMU_PX = 6350  # 12192000 EMU (13.333 in) / 1920 px

# font key -> (pptx family, css family, css weight)
FONTS = {
    "light": ("Manrope Light", "Manrope", 300),
    "regular": ("Manrope", "Manrope", 400),
    "semibold": ("Manrope SemiBold", "Manrope", 600),
    "mono": ("IBM Plex Mono", "IBM Plex Mono", 400),
}


class Slide:
    def __init__(self, sid, bg, title=""):
        self.sid = sid
        self.bg = bg
        self.title = title
        self.items = []

    # --- elements -------------------------------------------------------
    def img(self, src, x, y, w, h, fx=0.5, fy=0.5, alt=""):
        self.items.append(dict(kind="img", src=src, x=x, y=y, w=w, h=h, fx=fx, fy=fy, alt=alt))
        return self

    def rect(self, x, y, w, h, color, alpha=1.0):
        self.items.append(dict(kind="rect", x=x, y=y, w=w, h=h, color=color, alpha=alpha))
        return self

    def grad(self, x, y, w, h, angle, stops):
        """angle in CSS degrees (180 = top->bottom); stops: [(pos 0..1, hex, alpha)]."""
        self.items.append(dict(kind="grad", x=x, y=y, w=w, h=h, angle=angle, stops=stops))
        return self

    def text(self, x, y, w, h, paras, font="regular", size=30, color="#000000", lh=1.4,
             ls=0.0, align="l", valign="t", pgap=0, upper=False):
        if isinstance(paras, str):
            paras = [paras]
        if upper:
            paras = [p.upper() for p in paras]
        self.items.append(dict(kind="text", x=x, y=y, w=w, h=h, paras=paras, font=font, size=size,
                               color=color, lh=lh, ls=ls, align=align, valign=valign, pgap=pgap))
        return self


# ---------------------------------------------------------------------------
# geometry helpers


def cover_crop(src_w, src_h, box_w, box_h, fx, fy):
    """Fractions to crop (l, t, r, b) so the image covers the box like
    CSS object-fit:cover + object-position:fx fy."""
    src_ar = src_w / src_h
    box_ar = box_w / box_h
    if src_ar > box_ar:  # too wide: crop sides
        keep = box_ar / src_ar
        extra = 1 - keep
        return extra * fx, 0.0, extra * (1 - fx), 0.0
    keep = src_ar / box_ar
    extra = 1 - keep
    return 0.0, extra * fy, 0.0, extra * (1 - fy)


def hex_rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def rgba(h, a):
    h = h.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{a:g})"


# ---------------------------------------------------------------------------
# PPTX

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _strip_style(shape):
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _px(v):
    return Emu(int(round(v * EMU_PX)))


def _pptx_rect(slide, it):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _px(it["x"]), _px(it["y"]), _px(it["w"]), _px(it["h"]))
    _strip_style(shp)
    shp.fill.solid()
    shp.fill.fore_color.rgb = hex_rgb(it["color"])
    shp.line.fill.background()
    if it["alpha"] < 1:
        clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(it["alpha"] * 100000)))
    shp.name = "Scrim"


def _pptx_grad(slide, it):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _px(it["x"]), _px(it["y"]), _px(it["w"]), _px(it["h"]))
    _strip_style(shp)
    shp.line.fill.background()
    spPr = shp._element.spPr
    for old in spPr.findall(qn("a:solidFill")) + spPr.findall(qn("a:noFill")):
        spPr.remove(old)
    grad = etree.Element(qn("a:gradFill"))
    grad.set("rotWithShape", "1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, color, alpha in it["stops"]:
        gs = etree.SubElement(gs_lst, qn("a:gs"))
        gs.set("pos", str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", color.lstrip("#").upper())
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(int(((it["angle"] - 90) % 360) * 60000)))
    lin.set("scaled", "0")
    # gradFill must come right after the geometry element
    geom = spPr.find(qn("a:prstGeom"))
    geom.addnext(grad)
    shp.name = "Gradient"


def _pptx_img(slide, it, asset_dir):
    path = os.path.join(asset_dir, it["src"])
    with Image.open(path) as im:
        sw, sh = im.size
    pic = slide.shapes.add_picture(path, _px(it["x"]), _px(it["y"]), _px(it["w"]), _px(it["h"]))
    l, t, r, b = cover_crop(sw, sh, it["w"], it["h"], it["fx"], it["fy"])
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    pic.name = it["src"]
    if it.get("alt"):
        pic._element.nvPicPr.cNvPr.set("descr", it["alt"])


def _pptx_text(slide, it):
    tb = slide.shapes.add_textbox(_px(it["x"]), _px(it["y"]), _px(it["w"]), _px(it["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[it["valign"]]
    family = FONTS[it["font"]][0]
    size_pt = it["size"] * 0.5
    for i, para in enumerate(it["paras"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[it["align"]]
        p.line_spacing = Pt(size_pt * it["lh"])
        if i > 0 and it["pgap"]:
            p.space_before = Pt(it["pgap"] * 0.5)
        lines = para.split("\n")
        for j, line in enumerate(lines):
            if j > 0:
                p.add_line_break()
            run = p.add_run()
            run.text = line
            f = run.font
            f.name = family
            f.size = Pt(size_pt)
            f.color.rgb = hex_rgb(it["color"])
            if it["ls"]:
                run._r.get_or_add_rPr().set("spc", str(int(round(it["ls"] * 0.5 * 100))))
    tb.name = (it["paras"][0][:40] if it["paras"] else "Text")


def build_pptx(slides, asset_dir, out_path, notes=None):
    prs = Presentation()
    prs.slide_width = Emu(W * EMU_PX)
    prs.slide_height = Emu(H * EMU_PX)
    blank = prs.slide_layouts[6]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = hex_rgb(s.bg)
        for it in s.items:
            {"img": lambda it: _pptx_img(slide, it, asset_dir),
             "rect": lambda it: _pptx_rect(slide, it),
             "grad": lambda it: _pptx_grad(slide, it),
             "text": lambda it: _pptx_text(slide, it)}[it["kind"]](it)
        if notes and s.sid in notes:
            slide.notes_slide.notes_text_frame.text = notes[s.sid]
    prs.core_properties.title = "SONU SS27 — Director's Note"
    prs.save(out_path)


# ---------------------------------------------------------------------------
# HTML


def _html_item(it, asset_prefix):
    box = f"left:{it['x']}px;top:{it['y']}px;width:{it['w']}px;height:{it['h']}px"
    k = it["kind"]
    if k == "img":
        return (f'<img class="im" src="{asset_prefix}{html.escape(it["src"])}" alt="{html.escape(it.get("alt", ""))}" '
                f'style="{box};object-position:{it["fx"] * 100:g}% {it["fy"] * 100:g}%">')
    if k == "rect":
        return f'<div class="bx" style="{box};background:{rgba(it["color"], it["alpha"])}"></div>'
    if k == "grad":
        stops = ", ".join(f"{rgba(c, a)} {p * 100:g}%" for p, c, a in it["stops"])
        return f'<div class="bx" style="{box};background:linear-gradient({it["angle"]}deg, {stops})"></div>'
    if k == "text":
        _, fam, weight = FONTS[it["font"]]
        just = {"t": "flex-start", "m": "center", "b": "flex-end"}[it["valign"]]
        align = {"l": "left", "c": "center", "r": "right"}[it["align"]]
        style = (f"{box};font-family:'{fam}';font-weight:{weight};font-size:{it['size']}px;"
                 f"line-height:{it['lh']};letter-spacing:{it['ls']}px;color:{it['color']};"
                 f"text-align:{align};justify-content:{just}")
        paras = []
        for i, p in enumerate(it["paras"]):
            gap = f' style="margin-top:{it["pgap"]}px"' if i and it["pgap"] else ""
            paras.append(f"<p{gap}>" + "<br>".join(html.escape(l) for l in p.split("\n")) + "</p>")
        return f'<div class="tx" style="{style}">' + "".join(paras) + "</div>"
    raise ValueError(k)


def build_html(slides, template_path, out_path, asset_prefix, font_prefix):
    sections = []
    for i, s in enumerate(slides):
        body = "\n    ".join(_html_item(it, asset_prefix) for it in s.items)
        sections.append(f'  <section class="slide" id="{s.sid}" data-title="{html.escape(s.title)}" '
                        f'style="background:{s.bg}">\n    {body}\n  </section>')
    with open(template_path, encoding="utf8") as f:
        tpl = f.read()
    out = tpl.replace("{{SLIDES}}", "\n".join(sections)).replace("{{FONTS}}", font_prefix)
    with open(out_path, "w", encoding="utf8") as f:
        f.write(out)
